"""
Logica del inventario: altas, escaneos y reportes.

Aqui vive la regla de negocio; ni la ventana de escaneo ni el generador de
etiquetas saben como se actualiza la base, solo llaman a estas funciones.
"""

from __future__ import annotations

import csv
import re
import sqlite3
from dataclasses import dataclass, field
from pathlib import Path

import db

# Modos de operacion de la ventana de escaneo.
MODO_AUTO = "AUTO"        # articulo unico: alterna salida/entrada. Consumible: salida.
MODO_SALIDA = "SALIDA"    # siempre descuenta / marca como fuera
MODO_ENTRADA = "ENTRADA"  # siempre devuelve / marca como en bodega

# Tipos de articulo. En la base se guarda la clave; al usuario se le
# muestra el nombre largo (o el corto en las tablas).
#   UNICO      una pieza; existencia 0/1 solo indica si esta en la bodega.
#   AGRUPADO   varias piezas iguales con una sola etiqueta que se prestan y
#              se devuelven. existencia = piezas DISPONIBLES en bodega; lo
#              prestado vive en la tabla 'prestamos' (quien y cuantas).
#   CONSUMIBLE material que se gasta; existencia = lo que queda.
TIPOS = db.TIPOS
NOMBRE_TIPO = {
    "UNICO": "Herramienta única (1 pieza)",
    "AGRUPADO": "Herramienta con cantidad (varias iguales, se devuelven)",
    "CONSUMIBLE": "Material (se gasta)",
}
TIPO_CORTO = {"UNICO": "herram.", "AGRUPADO": "con cant.",
              "CONSUMIBLE": "material"}

# Margen para comparar cantidades con decimales (0.1 + 0.2 != 0.3).
_EPS = 1e-9

# Cuando alguien regresa piezas que tenia otra persona, la nota del
# movimiento lleva "[PER-0001]" con el gafete de quien las debia. Asi el
# historial basta para reconstruir quien debe que (ver reconstruir_prestamos).
_REGEX_DE_PERSONA = re.compile(r"regresa lo de .*\[([A-Z0-9]+-\d+)\]")


def fmt(numero: float) -> str:
    """Muestra 3 en vez de 3.0, pero conserva 2.5 cuando hace falta."""
    numero = float(numero)
    return str(int(numero)) if numero.is_integer() else f"{numero:g}"


def fmt_dinero(numero: float) -> str:
    """Formatea con separador de miles: 12345.5 -> '$12,345.50'."""
    return f"${float(numero or 0):,.2f}"


def _prestado(fila) -> float:
    """Piezas prestadas de un AGRUPADO, si la fila trae la columna 'prestado'
    (la agregan listar_articulos, buscar_articulo y bajo_stock)."""
    return float(fila["prestado"] or 0) if "prestado" in fila.keys() else 0.0


def valor_articulo(fila) -> float:
    """Valor monetario de un articulo, este o no en la bodega.

    UNICO vale 1 pieza; CONSUMIBLE lo que queda; AGRUPADO lo disponible mas
    lo prestado (prestar no cambia su valor total).
    """
    if fila["tipo"] == "UNICO":
        cantidad = 1
    elif fila["tipo"] == "AGRUPADO":
        cantidad = float(fila["existencia"] or 0) + _prestado(fila)
    else:
        cantidad = fila["existencia"]
    return float(fila["valor_unitario"] or 0) * float(cantidad or 0)


def texto_existencia(fila) -> str:
    """Lo que se muestra en la columna 'Exist.' de las listas."""
    if fila["tipo"] == "AGRUPADO":
        return f"{fmt(fila['existencia'])} disp. / {fmt(_prestado(fila))} fuera"
    return f"{fmt(fila['existencia'])} {fila['unidad']}".strip()


def tiene_tabla_prestamos(con: sqlite3.Connection) -> bool:
    """False en fotos publicadas por versiones anteriores (visor)."""
    return con.execute(
        "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = 'prestamos'"
    ).fetchone() is not None


def _sql_prestado(con, alias: str = "a") -> str:
    """Subconsulta con las piezas prestadas del articulo `alias`.codigo.
    En una foto vieja sin tabla 'prestamos' vale 0."""
    if not tiene_tabla_prestamos(con):
        return "0"
    return (f"(SELECT COALESCE(SUM(pr.cantidad), 0) FROM prestamos pr "
            f"WHERE pr.codigo = {alias}.codigo)")


@dataclass
class Resultado:
    """Lo que devuelve un escaneo, listo para pintarse en pantalla."""
    ok: bool
    nivel: str          # 'ok' | 'aviso' | 'error' | 'persona'
    titulo: str
    detalle: str = ""
    articulo: dict = field(default_factory=dict)
    persona: dict = field(default_factory=dict)


# --------------------------------------------------------------------------
# Consultas basicas
# --------------------------------------------------------------------------

def buscar_articulo(con: sqlite3.Connection, codigo: str) -> sqlite3.Row | None:
    return con.execute(
        f"SELECT a.*, {_sql_prestado(con)} AS prestado "
        f"FROM articulos a WHERE a.codigo = ?", (codigo,)
    ).fetchone()


def buscar_persona(con: sqlite3.Connection, codigo: str) -> sqlite3.Row | None:
    return con.execute(
        "SELECT * FROM personas WHERE codigo = ? AND activo = 1", (codigo,)
    ).fetchone()


def listar_articulos(con: sqlite3.Connection, filtro: str = "",
                     incluir_bajas: bool = False, ubicacion: str = "",
                     estado: str = "") -> list[sqlite3.Row]:
    condiciones = [] if incluir_bajas else ["activo = 1"]
    parametros: list = []
    if filtro:
        patron = f"%{filtro}%"
        condiciones.append(
            "(codigo LIKE ? OR nombre LIKE ? OR categoria LIKE ? "
            " OR ubicacion LIKE ? OR marca LIKE ? OR modelo LIKE ? "
            " OR numero_serie LIKE ?)"
        )
        parametros += [patron] * 7
    if ubicacion:
        condiciones.append("ubicacion = ?")
        parametros.append(ubicacion)
    prestado = _sql_prestado(con)
    if estado == "FUERA":
        # Un AGRUPADO esta "fuera" si tiene al menos una pieza prestada,
        # aunque le queden otras en la bodega.
        condiciones.append(
            f"(estado = 'FUERA' OR (tipo = 'AGRUPADO' AND {prestado} > 0))")
    elif estado:
        condiciones.append("estado = ?")
        parametros.append(estado)
    donde = f"WHERE {' AND '.join(condiciones)}" if condiciones else ""
    return con.execute(
        f"SELECT a.*, {prestado} AS prestado FROM articulos a {donde} "
        f"ORDER BY codigo", parametros
    ).fetchall()


def listar_personas(con: sqlite3.Connection,
                    incluir_bajas: bool = False) -> list[sqlite3.Row]:
    donde = "" if incluir_bajas else "WHERE activo = 1"
    return con.execute(
        f"SELECT * FROM personas {donde} ORDER BY codigo"
    ).fetchall()


def ultimos_movimientos(con: sqlite3.Connection, limite: int = 25,
                        filtro: str = "") -> list[sqlite3.Row]:
    condicion, parametros = "", []
    if filtro:
        patron = f"%{filtro}%"
        condicion = ("WHERE m.codigo LIKE ? OR a.nombre LIKE ? "
                    "OR p.nombre LIKE ?")
        parametros = [patron, patron, patron]
    return con.execute(
        f"""
        SELECT m.*, a.nombre AS articulo, a.unidad, p.nombre AS persona_nombre
        FROM movimientos m
        LEFT JOIN articulos a ON a.codigo = m.codigo
        LEFT JOIN personas  p ON p.codigo = m.persona
        {condicion}
        ORDER BY m.id DESC LIMIT ?
        """,
        (*parametros, limite),
    ).fetchall()


def pendientes(con: sqlite3.Connection, dias: int = 0) -> list[sqlite3.Row]:
    """Articulos que estan fuera de la bodega, con quien los saco y desde cuando.

    Un UNICO (o consumible agotado) sale en una sola fila. Un AGRUPADO sale
    una vez por cada persona que tiene piezas, con su 'cantidad'. Todas las
    filas traen 'valor' = lo que vale lo que esta fuera en esa linea.
    """
    sql = """
        SELECT a.codigo AS codigo, a.nombre, a.tipo, a.existencia, a.unidad,
               a.valor_unitario,
               CASE WHEN a.tipo = 'UNICO' THEN 1 ELSE a.existencia END AS cantidad,
               a.valor_unitario * (CASE WHEN a.tipo = 'UNICO' THEN 1
                                   ELSE a.existencia END) AS valor,
               m.persona AS persona,
               m.fecha AS desde, COALESCE(p.nombre, '(sin registrar)') AS quien,
               CAST(julianday('now', 'localtime') - julianday(m.fecha) AS INTEGER) AS dias
        FROM articulos a
        LEFT JOIN movimientos m ON m.id = (
            SELECT id FROM movimientos
            WHERE codigo = a.codigo AND tipo = 'SALIDA'
            ORDER BY id DESC LIMIT 1
        )
        LEFT JOIN personas p ON p.codigo = m.persona
        WHERE a.estado = 'FUERA' AND a.activo = 1 AND a.tipo <> 'AGRUPADO'
          AND CAST(julianday('now','localtime') - julianday(COALESCE(m.fecha,
              datetime('now','localtime'))) AS INTEGER) >= ?
    """
    parametros: list = [dias]
    if tiene_tabla_prestamos(con):
        sql += """
        UNION ALL
        SELECT a.codigo AS codigo, a.nombre, a.tipo, a.existencia, a.unidad,
               a.valor_unitario,
               pr.cantidad AS cantidad,
               a.valor_unitario * pr.cantidad AS valor,
               pr.persona AS persona,
               pr.desde AS desde, COALESCE(p.nombre, pr.persona) AS quien,
               CAST(julianday('now', 'localtime') - julianday(pr.desde) AS INTEGER) AS dias
        FROM prestamos pr
        JOIN articulos a ON a.codigo = pr.codigo
        LEFT JOIN personas p ON p.codigo = pr.persona
        WHERE a.activo = 1 AND a.tipo = 'AGRUPADO'
          AND CAST(julianday('now','localtime') - julianday(pr.desde)
              AS INTEGER) >= ?
        """
        parametros.append(dias)
    sql += " ORDER BY dias DESC, codigo"
    return con.execute(sql, parametros).fetchall()


def prestamos_de(con: sqlite3.Connection, codigo: str) -> list[sqlite3.Row]:
    """Quien tiene piezas prestadas de un AGRUPADO y cuantas."""
    if not tiene_tabla_prestamos(con):
        return []
    return con.execute(
        "SELECT pr.persona, COALESCE(p.nombre, pr.persona) AS nombre, "
        "       pr.cantidad, pr.desde "
        "FROM prestamos pr LEFT JOIN personas p ON p.codigo = pr.persona "
        "WHERE pr.codigo = ? ORDER BY pr.desde, p.nombre",
        (codigo.strip().upper(),),
    ).fetchall()


def prestamos_activos(con: sqlite3.Connection) -> list[sqlite3.Row]:
    """Todos los prestamos de AGRUPADO, para exportar."""
    if not tiene_tabla_prestamos(con):
        return []
    return con.execute(
        "SELECT pr.codigo, a.nombre AS articulo, pr.persona, "
        "       COALESCE(p.nombre, '') AS nombre_persona, pr.cantidad, "
        "       a.unidad, a.valor_unitario, "
        "       a.valor_unitario * pr.cantidad AS valor, pr.desde "
        "FROM prestamos pr "
        "JOIN articulos a ON a.codigo = pr.codigo "
        "LEFT JOIN personas p ON p.codigo = pr.persona "
        "ORDER BY pr.codigo, pr.desde"
    ).fetchall()


def bajo_stock(con: sqlite3.Connection) -> list[sqlite3.Row]:
    """Consumibles y herramientas con cantidad cuya existencia DISPONIBLE
    esta en o bajo su minimo. Solo avisa si el minimo es mayor que 0, que
    es el valor por omision, asi que un AGRUPADO no genera avisos a menos
    que alguien le ponga un minimo a proposito."""
    return con.execute(
        f"SELECT a.*, {_sql_prestado(con)} AS prestado FROM articulos a "
        "WHERE tipo IN ('CONSUMIBLE', 'AGRUPADO') AND activo = 1 "
        "AND minimo > 0 AND existencia <= minimo "
        "ORDER BY existencia / NULLIF(minimo, 0)"
    ).fetchall()


def _sql_cantidades(con) -> tuple[str, str]:
    """(en bodega, prestado) por articulo, en piezas, para los SUM de valor.

    - UNICO: su 'existencia' se pone en 0 al prestarse (solo marca
      presencia), asi que cuenta 1 pieza en bodega o 1 fuera segun su estado.
    - CONSUMIBLE: lo que queda en existencia; lo que salio se gasto.
    - AGRUPADO: existencia = disponibles en bodega; prestado = suma de la
      tabla prestamos. El total (la suma) no cambia al prestar.
    """
    prestado = _sql_prestado(con)
    en_bodega = ("(CASE WHEN tipo = 'UNICO' THEN (estado = 'EN_BODEGA')"
                 " WHEN tipo = 'AGRUPADO' THEN existencia"
                 " WHEN estado = 'EN_BODEGA' THEN existencia ELSE 0 END)")
    fuera = ("(CASE WHEN tipo = 'UNICO' THEN (estado = 'FUERA')"
             f" WHEN tipo = 'AGRUPADO' THEN {prestado}"
             " WHEN estado = 'FUERA' THEN existencia ELSE 0 END)")
    return en_bodega, fuera


def valor_inventario(con: sqlite3.Connection) -> dict:
    """Valor del inventario activo, valuado a costo unitario.

    Devuelve en_bodega ("Valor en bodega"), fuera ("Prestado") y
    total_bodega ("Total" = en_bodega + fuera; tambien como 'total').
    Una herramienta prestada sigue valiendo lo mismo: solo cambia de columna.
    """
    en_bodega, fuera = _sql_cantidades(con)
    fila = con.execute(
        f"SELECT COALESCE(SUM(valor_unitario * {en_bodega}), 0) AS en_bodega,"
        f" COALESCE(SUM(valor_unitario * {fuera}), 0) AS fuera,"
        f" COUNT(*) AS articulos"
        f" FROM articulos a WHERE activo = 1"
    ).fetchone()
    resultado = dict(fila)
    resultado["total_bodega"] = resultado["en_bodega"] + resultado["fuera"]
    resultado["total"] = resultado["total_bodega"]
    return resultado


def valor_por_categoria(con: sqlite3.Connection) -> list[sqlite3.Row]:
    en_bodega, fuera = _sql_cantidades(con)
    return con.execute(
        f"SELECT COALESCE(NULLIF(categoria, ''), '(sin categoria)') AS categoria,"
        f" COUNT(*) AS articulos,"
        f" SUM(valor_unitario * ({en_bodega} + {fuera})) AS valor"
        f" FROM articulos a WHERE activo = 1"
        f" GROUP BY categoria ORDER BY valor DESC"
    ).fetchall()


# --------------------------------------------------------------------------
# Altas
# --------------------------------------------------------------------------

def agregar_articulo(
    con: sqlite3.Connection,
    nombre: str,
    tipo: str = "UNICO",
    categoria: str = "",
    ubicacion: str = "",
    unidad: str = "pza",
    existencia: float | None = None,
    minimo: float = 0,
    marca: str = "",
    modelo: str = "",
    numero_serie: str = "",
    valor_unitario: float = 0,
    notas: str = "",
    activo: bool = True,
    codigo: str | None = None,
) -> str:
    tipo = tipo.strip().upper()
    if tipo not in TIPOS:
        raise ValueError("El tipo debe ser UNICO, AGRUPADO o CONSUMIBLE")

    if existencia is None:
        existencia = 1 if tipo == "UNICO" else 0
    if tipo == "UNICO":
        existencia = 1 if existencia else 0
    _validar_no_negativos(existencia=existencia, minimo=minimo,
                          valor_unitario=valor_unitario)

    if codigo is None:
        # Las herramientas con cantidad son herramienta: tambien llevan HER-.
        prefijo = (db.PREFIJO_CONSUMIBLE if tipo == "CONSUMIBLE"
                   else db.PREFIJO_UNICO)
        codigo = db.siguiente_codigo(con, prefijo)
    codigo = codigo.strip().upper()

    con.execute(
        "INSERT INTO articulos "
        "(codigo, nombre, categoria, tipo, ubicacion, unidad, existencia, "
        " minimo, marca, modelo, numero_serie, valor_unitario, estado, "
        " notas, activo) "
        "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (codigo, nombre.strip(), categoria.strip(), tipo, ubicacion.strip(),
         unidad.strip() or "pza", float(existencia), float(minimo),
         marca.strip(), modelo.strip(), numero_serie.strip(),
         float(valor_unitario or 0),
         "EN_BODEGA" if existencia > 0 else "FUERA", notas.strip(),
         1 if activo else 0),
    )
    con.commit()
    return codigo


def agregar_persona(con: sqlite3.Connection, nombre: str,
                    codigo: str | None = None) -> str:
    codigo = (codigo or db.siguiente_codigo(con, db.PREFIJO_PERSONA)).strip().upper()
    con.execute("INSERT INTO personas (codigo, nombre) VALUES (?, ?)",
                (codigo, nombre.strip()))
    con.commit()
    return codigo


CAMPOS_EDITABLES = ("nombre", "categoria", "tipo", "ubicacion", "unidad",
                    "existencia", "minimo", "marca", "modelo",
                    "numero_serie", "valor_unitario", "notas")


def _validar_no_negativos(**valores) -> None:
    for nombre, valor in valores.items():
        if valor is not None and float(valor) < 0:
            raise ValueError(f"El campo '{nombre}' no puede ser negativo")


def actualizar_articulo(con: sqlite3.Connection, codigo: str, **campos) -> None:
    """Modifica un articulo ya existente. Solo toca los campos que le pasas.

    Si cambia el tipo a AGRUPADO, reconstruye sus prestamos a partir del
    historial (sin tocar la existencia). Si deja de ser AGRUPADO, exige que
    no tenga piezas prestadas (lanza ValueError). Conviene mostrar antes
    resumen_cambio_tipo() al usuario.
    """
    codigo = codigo.strip().upper()
    cambios = {k: v for k, v in campos.items() if k in CAMPOS_EDITABLES}
    if not cambios:
        return

    actual = buscar_articulo(con, codigo)
    if "tipo" in cambios:
        cambios["tipo"] = str(cambios["tipo"]).strip().upper()
        if cambios["tipo"] not in TIPOS:
            raise ValueError("El tipo debe ser UNICO, AGRUPADO o CONSUMIBLE")
    tipo_antes = actual["tipo"] if actual else None
    tipo_nuevo = cambios.get("tipo", tipo_antes)

    if tipo_antes == "AGRUPADO" and tipo_nuevo != "AGRUPADO" and prestamos_de(con, codigo):
        raise ValueError(
            "No se puede cambiar el tipo: todavia hay piezas prestadas "
            f"({_quienes_lo_tienen(con, codigo)}). Registra primero su devolucion.")
    nuevos_prestamos = (reconstruir_prestamos(con, codigo)
                        if tipo_nuevo == "AGRUPADO" and tipo_antes != "AGRUPADO"
                        else None)

    if cambios.get("tipo") == "UNICO":
        cambios["existencia"] = 1 if (actual and actual["estado"] == "EN_BODEGA") else 0
        cambios["minimo"] = 0

    _validar_no_negativos(existencia=cambios.get("existencia"),
                          minimo=cambios.get("minimo"),
                          valor_unitario=cambios.get("valor_unitario"))

    if "existencia" in cambios:
        cambios["estado"] = "EN_BODEGA" if float(cambios["existencia"]) > 0 else "FUERA"

    asignaciones = ", ".join(f"{campo} = ?" for campo in cambios)
    with con:  # una sola transaccion: el cambio de tipo y sus prestamos
        con.execute(f"UPDATE articulos SET {asignaciones} WHERE codigo = ?",
                    (*cambios.values(), codigo))
        if nuevos_prestamos is not None:
            con.execute("DELETE FROM prestamos WHERE codigo = ?", (codigo,))
            con.executemany(
                "INSERT INTO prestamos (codigo, persona, cantidad, desde) "
                "VALUES (?,?,?,?)",
                [(codigo, p["persona"], p["cantidad"], p["desde"])
                 for p in nuevos_prestamos])


def reconstruir_prestamos(con: sqlite3.Connection, codigo: str) -> list[dict]:
    """Calcula quien deberia tener piezas de un articulo segun su historial.

    Se usa al convertir un articulo a AGRUPADO. No escribe nada.
    - UNICO: si esta fuera, 1 pieza a quien hizo la ultima SALIDA (sumar
      salidas menos entradas por persona fallaria cuando la regresa otro).
    - Otros: por persona, SALIDAs menos ENTRADAs; solo cuenta si queda > 0.
      Una ENTRADA marcada "regresa lo de ... [PER-xxxx]" se descuenta a esa
      persona, no a quien la registro.
    Devuelve [{persona, nombre, cantidad, desde}].
    """
    codigo = codigo.strip().upper()
    art = buscar_articulo(con, codigo)
    if art is None:
        return []

    if art["tipo"] == "UNICO":
        if art["estado"] != "FUERA":
            return []
        ultima = con.execute(
            "SELECT persona, fecha FROM movimientos WHERE codigo = ? "
            "AND tipo = 'SALIDA' ORDER BY id DESC LIMIT 1", (codigo,)).fetchone()
        if ultima is None or not ultima["persona"]:
            return []
        return [{"persona": ultima["persona"],
                 "nombre": _nombre_persona(con, ultima["persona"]),
                 "cantidad": 1.0, "desde": ultima["fecha"]}]

    saldo: dict[str, float] = {}
    desde: dict[str, str] = {}
    for mov in con.execute(
            "SELECT persona, tipo, cantidad, fecha, nota FROM movimientos "
            "WHERE codigo = ? AND tipo IN ('SALIDA', 'ENTRADA') ORDER BY id",
            (codigo,)):
        quien = mov["persona"]
        if mov["tipo"] == "ENTRADA":
            marca = _REGEX_DE_PERSONA.search(mov["nota"] or "")
            if marca:
                quien = marca.group(1)
        if not quien:
            continue
        antes = saldo.get(quien, 0.0)
        cambio = float(mov["cantidad"]) * (1 if mov["tipo"] == "SALIDA" else -1)
        saldo[quien] = antes + cambio
        if antes <= _EPS < saldo[quien]:
            desde[quien] = mov["fecha"]  # desde cuando debe piezas

    return [{"persona": quien, "nombre": _nombre_persona(con, quien),
             "cantidad": cantidad, "desde": desde[quien]}
            for quien, cantidad in sorted(saldo.items())
            if cantidad > _EPS and buscar_persona_cualquiera(con, quien)]


def buscar_persona_cualquiera(con, codigo) -> sqlite3.Row | None:
    """Como buscar_persona, pero tambien encuentra a las dadas de baja."""
    return con.execute("SELECT * FROM personas WHERE codigo = ?",
                       (codigo,)).fetchone()


def resumen_cambio_tipo(con: sqlite3.Connection, codigo: str,
                        tipo_nuevo: str) -> tuple[bool, str, bool]:
    """Lo que pasaria al cambiar el tipo de un articulo, para avisar antes.

    Devuelve (permitido, mensaje, pedir_confirmacion).
    """
    art = buscar_articulo(con, codigo.strip().upper())
    tipo_nuevo = tipo_nuevo.strip().upper()
    if art is None or art["tipo"] == tipo_nuevo:
        return True, "", False
    unidad = art["unidad"] or "pza"

    if art["tipo"] == "AGRUPADO":
        if prestamos_de(con, art["codigo"]):
            return (False,
                    "No se puede cambiar el tipo mientras haya piezas "
                    f"prestadas.\n\nLas tienen: {_quienes_lo_tienen(con, art['codigo'])}"
                    "\n\nRegistra primero su devolucion.", False)
        if tipo_nuevo == "UNICO" and float(art["existencia"]) > 1:
            return (True,
                    f"Hay {fmt(art['existencia'])} {unidad} disponibles. Como "
                    f"herramienta unica contara como UNA sola pieza.\n\n"
                    "¿Continuar?", True)
        return True, "", False

    if tipo_nuevo != "AGRUPADO":
        return True, "", False

    hay_movimientos = con.execute(
        "SELECT 1 FROM movimientos WHERE codigo = ? LIMIT 1",
        (art["codigo"],)).fetchone() is not None
    if not hay_movimientos:
        return True, "", False

    prestamos = reconstruir_prestamos(con, art["codigo"])
    disponibles = float(art["existencia"])
    total_prestado = sum(p["cantidad"] for p in prestamos)
    lineas = [f"{art['nombre']} pasara a ser:\n{NOMBRE_TIPO['AGRUPADO']}.", "",
              f"Disponibles en bodega: {fmt(disponibles)} {unidad} (no cambia)"]
    if prestamos:
        lineas.append("Prestamos reconstruidos desde el historial "
                      "(salidas menos entradas de cada persona):")
        lineas += [f"   - {p['nombre']}: {fmt(p['cantidad'])} {unidad}"
                   for p in prestamos]
    else:
        lineas.append("El historial no muestra piezas pendientes de devolver.")
    lineas += ["", f"Total: {fmt(disponibles)} + {fmt(total_prestado)} = "
                   f"{fmt(disponibles + total_prestado)} {unidad}", "",
               "¿Continuar?"]
    return True, "\n".join(lineas), True


def baja_articulo(con: sqlite3.Connection, codigo: str, activo: bool = False) -> None:
    """Da de baja (o reactiva) un articulo sin borrar su historial."""
    con.execute("UPDATE articulos SET activo = ? WHERE codigo = ?",
                (1 if activo else 0, codigo.strip().upper()))
    con.commit()


def actualizar_persona(con: sqlite3.Connection, codigo: str, nombre: str) -> None:
    con.execute("UPDATE personas SET nombre = ? WHERE codigo = ?",
                (nombre.strip(), codigo.strip().upper()))
    con.commit()


def baja_persona(con: sqlite3.Connection, codigo: str, activo: bool = False) -> None:
    con.execute("UPDATE personas SET activo = ? WHERE codigo = ?",
                (1 if activo else 0, codigo.strip().upper()))
    con.commit()


def tiene_pendientes(con: sqlite3.Connection, codigo_persona: str) -> list[sqlite3.Row]:
    """Herramientas que siguen fuera a nombre de esta persona (incluye las
    piezas de herramientas con cantidad que tiene prestadas)."""
    codigo_persona = codigo_persona.strip().upper()
    sql = """
        SELECT a.codigo, a.nombre FROM articulos a
        WHERE a.estado = 'FUERA' AND a.activo = 1 AND a.tipo <> 'AGRUPADO' AND (
            SELECT m.persona FROM movimientos m
            WHERE m.codigo = a.codigo AND m.tipo = 'SALIDA'
            ORDER BY m.id DESC LIMIT 1) = ?
        """
    parametros = [codigo_persona]
    if tiene_tabla_prestamos(con):
        sql += """
        UNION ALL
        SELECT a.codigo, a.nombre || ' (' || printf('%g', pr.cantidad) || ')'
        FROM prestamos pr
        JOIN articulos a ON a.codigo = pr.codigo
        WHERE a.activo = 1 AND pr.persona = ?
        """
        parametros.append(codigo_persona)
    return con.execute(sql, parametros).fetchall()


def importar_csv(con: sqlite3.Connection, ruta: Path) -> tuple[int, list[str]]:
    """Carga articulos desde un CSV. Columnas reconocidas:
    nombre, tipo, categoria, ubicacion, unidad, existencia, minimo,
    marca, modelo, numero_serie, valor_unitario, notas, activo, codigo
    Solo 'nombre' es obligatoria.
    """
    altas, errores = 0, []
    with open(ruta, newline="", encoding="utf-8-sig") as archivo:
        for num, fila in enumerate(csv.DictReader(archivo), start=2):
            fila = {(k or "").strip().lower(): (v or "").strip()
                    for k, v in fila.items()}
            if not fila.get("nombre"):
                continue
            try:
                activo_txt = fila.get("activo", "1").strip().lower()
                agregar_articulo(
                    con,
                    nombre=fila["nombre"],
                    tipo=fila.get("tipo") or "UNICO",
                    categoria=fila.get("categoria", ""),
                    ubicacion=fila.get("ubicacion", ""),
                    unidad=fila.get("unidad") or "pza",
                    existencia=float(fila["existencia"]) if fila.get("existencia") else None,
                    minimo=float(fila["minimo"]) if fila.get("minimo") else 0,
                    marca=fila.get("marca", ""),
                    modelo=fila.get("modelo", ""),
                    numero_serie=fila.get("numero_serie", ""),
                    valor_unitario=float(fila["valor_unitario"]) if fila.get("valor_unitario") else 0,
                    notas=fila.get("notas", ""),
                    activo=activo_txt not in ("0", "false", "no"),
                    codigo=fila.get("codigo") or None,
                )
                altas += 1
            except Exception as err:  # noqa: BLE001
                errores.append(f"linea {num}: {err}")
    return altas, errores


# --------------------------------------------------------------------------
# Sección de escaneo
# --------------------------------------------------------------------------

def registrar_escaneo(
    con: sqlite3.Connection,
    codigo: str,
    persona: str | None = None,
    modo: str = MODO_AUTO,
    cantidad: float = 1,
    nota: str = "",
    de_persona: str | None = None,
) -> Resultado:
    """Procesa un codigo leido por el escaner y actualiza la base.

    Para un AGRUPADO, MODO_ENTRADA es regresar piezas y cualquier otro modo
    es sacarlas. `de_persona` es de quien son las piezas que se regresan
    (por omision, del mismo operador); el movimiento queda a nombre del
    operador que las trae.
    """
    codigo = (codigo or "").strip().upper()
    if not codigo:
        return Resultado(False, "error", "Codigo vacio")

    # Gafete de persona: no mueve inventario, solo cambia quien esta activo.
    quien = buscar_persona(con, codigo)
    if quien:
        return Resultado(True, "persona", f"Operador: {quien['nombre']}",
                         "Escanea ahora las herramientas o materiales.",
                         persona=dict(quien))

    # Si el código no es un gafete, verificamos obligatoriamente que ya 
    # haya un responsable activo en la sesión antes de continuar.
    if not persona:
        return Resultado(False, "error", "Operador no seleccionado",
                         "Escanea primero tu gafete para registrar el movimiento.")

    # Articulos.
    art = buscar_articulo(con, codigo)
    if art is None:
        return Resultado(False, "error", f"{codigo} no esta registrado",
                         "Dalo de alta en el panel de gestion (boton Catalogo).")

    if not art["activo"]:
        return Resultado(False, "error", f"{art['nombre']} esta dado de baja",
                         "Reactivalo en el panel de gestion si vuelve a usarse.",
                         dict(art))

    cantidad = float(cantidad)
    if cantidad <= 0:
        return Resultado(False, "error", "La cantidad debe ser mayor que cero")

    if art["tipo"] == "UNICO":
        return _escanear_unico(con, art, persona, modo, nota)
    if art["tipo"] == "AGRUPADO":
        accion = "ENTRADA" if modo == MODO_ENTRADA else "SALIDA"
        return _escanear_agrupado(con, art, persona, cantidad, accion,
                                  de_persona=de_persona, nota=nota)
    return _escanear_consumible(con, art, persona, modo, cantidad, nota)


def _guardar_movimiento(con, codigo, persona, tipo, cantidad, nota) -> None:
    con.execute(
        "INSERT INTO movimientos (codigo, persona, tipo, cantidad, nota) "
        "VALUES (?,?,?,?,?)",
        (codigo, persona, tipo, float(cantidad), nota),
    )


def _escanear_unico(con, art, persona, modo, nota) -> Resultado:
    esta_dentro = art["estado"] == "EN_BODEGA"

    if modo == MODO_AUTO:
        accion = "SALIDA" if esta_dentro else "ENTRADA"
    else:
        accion = modo
        if accion == "SALIDA" and not esta_dentro:
            return Resultado(False, "aviso", f"{art['nombre']} ya estaba fuera",
                             _quien_lo_tiene(con, art["codigo"]), dict(art))
        if accion == "ENTRADA" and esta_dentro:
            return Resultado(False, "aviso", f"{art['nombre']} ya estaba en bodega",
                             "", dict(art))

    nuevo_estado = "FUERA" if accion == "SALIDA" else "EN_BODEGA"
    con.execute(
        "UPDATE articulos SET estado = ?, existencia = ? WHERE codigo = ?",
        (nuevo_estado, 0 if accion == "SALIDA" else 1, art["codigo"]),
    )
    _guardar_movimiento(con, art["codigo"], persona, accion, 1, nota)
    con.commit()

    if accion == "SALIDA":
        detalle = "Registrado como prestado" + (
            f" a {_nombre_persona(con, persona)}" if persona else " (sin operador)")
    else:
        detalle = f"De vuelta en su lugar: {art['ubicacion'] or 'sin ubicacion'}"

    return Resultado(True, "ok", f"{accion}: {art['nombre']}", detalle, dict(art))


def _escanear_consumible(con, art, persona, modo, cantidad, nota) -> Resultado:
    accion = "ENTRADA" if modo == MODO_ENTRADA else "SALIDA"
    existencia = float(art["existencia"])

    if accion == "SALIDA":
        if existencia <= 0:
            return Resultado(False, "error", f"Sin existencia: {art['nombre']}",
                             "Registra una entrada cuando llegue material.",
                             dict(art))
        if cantidad > existencia:
            return Resultado(False, "error",
                             f"Solo hay {fmt(existencia)} {art['unidad']} de "
                             f"{art['nombre']}",
                             f"Intentaste sacar {fmt(cantidad)}.", dict(art))
        nueva = existencia - cantidad
    else:
        nueva = existencia + cantidad

    con.execute(
        "UPDATE articulos SET existencia = ?, estado = ? WHERE codigo = ?",
        (nueva, "EN_BODEGA" if nueva > 0 else "FUERA", art["codigo"]),
    )
    _guardar_movimiento(con, art["codigo"], persona, accion, cantidad, nota)
    con.commit()

    nivel = "ok"
    detalle = f"Quedan {fmt(nueva)} {art['unidad']}"
    if art["minimo"] and nueva <= float(art["minimo"]):
        nivel = "aviso"
        detalle += f"  |  bajo el minimo de {fmt(art['minimo'])}, hay que resurtir"

    titulo = (f"{accion}: {fmt(cantidad)} {art['unidad']} de {art['nombre']}")
    return Resultado(True, nivel, titulo, detalle, dict(art))


def _escanear_agrupado(con, art, operador, cantidad, accion,
                       de_persona=None, nota="") -> Resultado:
    """Saca o regresa piezas de una herramienta con cantidad (AGRUPADO).

    SALIDA: resta de las disponibles y se las apunta al operador.
    ENTRADA: suma a las disponibles y se las descuenta a `de_persona` (por
    omision el mismo operador); si llega a 0 se borra su fila de prestamos.
    Todo va en una sola transaccion.
    """
    codigo, nombre = art["codigo"], art["nombre"]
    unidad = art["unidad"] or "pza"
    cantidad = float(cantidad)
    if cantidad <= 0:
        return Resultado(False, "error", "La cantidad debe ser mayor que cero")

    # Se relee dentro de la funcion por si cambio desde que se busco.
    existencia = float(con.execute(
        "SELECT existencia FROM articulos WHERE codigo = ?",
        (codigo,)).fetchone()["existencia"])

    if accion == "SALIDA":
        if existencia <= _EPS:
            quienes = _quienes_lo_tienen(con, codigo)
            return Resultado(False, "error", f"No quedan piezas de {nombre}",
                             f"Todas estan prestadas: {quienes}" if quienes
                             else "Registra un ajuste si llegaron piezas.",
                             dict(art))
        if cantidad > existencia + _EPS:
            return Resultado(False, "error",
                             f"Solo hay {fmt(existencia)} {unidad} disponibles "
                             f"de {nombre}",
                             f"Intentaste sacar {fmt(cantidad)}.", dict(art))
        nueva = existencia - cantidad
        with con:
            con.execute(
                "UPDATE articulos SET existencia = ?, estado = ? WHERE codigo = ?",
                (nueva, "EN_BODEGA" if nueva > _EPS else "FUERA", codigo))
            con.execute(
                "INSERT INTO prestamos (codigo, persona, cantidad) VALUES (?,?,?) "
                "ON CONFLICT (codigo, persona) "
                "DO UPDATE SET cantidad = cantidad + excluded.cantidad",
                (codigo, operador, cantidad))
            _guardar_movimiento(con, codigo, operador, "SALIDA", cantidad, nota)
        tiene = con.execute(
            "SELECT cantidad FROM prestamos WHERE codigo = ? AND persona = ?",
            (codigo, operador)).fetchone()["cantidad"]
        nivel = "ok"
        detalle = (f"Prestado a {_nombre_persona(con, operador)} "
                   f"(tiene {fmt(tiene)})  |  quedan {fmt(nueva)} {unidad} disponibles")
        if art["minimo"] and nueva <= float(art["minimo"]):
            nivel = "aviso"
            detalle += f"  |  en o bajo el minimo de {fmt(art['minimo'])}"
        return Resultado(True, nivel,
                         f"SALIDA: {fmt(cantidad)} {unidad} de {nombre}",
                         detalle, dict(art))

    # ENTRADA (devolucion)
    de = (de_persona or operador or "").strip().upper()
    fila = con.execute(
        "SELECT cantidad FROM prestamos WHERE codigo = ? AND persona = ?",
        (codigo, de)).fetchone()
    if fila is None:
        quienes = _quienes_lo_tienen(con, codigo)
        return Resultado(False, "error",
                         f"{_nombre_persona(con, de)} no tiene {nombre} prestado",
                         f"Lo tienen: {quienes}" if quienes
                         else "Nadie lo tiene prestado: todas las piezas estan "
                              "en la bodega.", dict(art))
    debe = float(fila["cantidad"])
    if cantidad > debe + _EPS:
        return Resultado(False, "error",
                         f"{_nombre_persona(con, de)} solo tiene {fmt(debe)} "
                         f"{unidad} de {nombre}",
                         f"Intentaste regresar {fmt(cantidad)}.", dict(art))

    if de != operador:
        marca = f"regresa lo de {_nombre_persona(con, de)} [{de}]"
        nota = f"{nota} · {marca}" if nota else marca
    restante = debe - cantidad
    nueva = existencia + cantidad
    with con:
        if restante <= _EPS:
            con.execute("DELETE FROM prestamos WHERE codigo = ? AND persona = ?",
                        (codigo, de))
        else:
            con.execute("UPDATE prestamos SET cantidad = ? "
                        "WHERE codigo = ? AND persona = ?", (restante, codigo, de))
        con.execute(
            "UPDATE articulos SET existencia = ?, estado = 'EN_BODEGA' "
            "WHERE codigo = ?", (nueva, codigo))
        _guardar_movimiento(con, codigo, operador, "ENTRADA", cantidad, nota)

    detalle = f"Disponibles: {fmt(nueva)} {unidad}"
    if de != operador:
        detalle += f"  |  regresadas por {_nombre_persona(con, operador)}"
    detalle += (f"  |  a {_nombre_persona(con, de)} le quedan {fmt(restante)}"
                if restante > _EPS else
                f"  |  {_nombre_persona(con, de)} ya no debe piezas")
    return Resultado(True, "ok",
                     f"ENTRADA: {fmt(cantidad)} {unidad} de {nombre}",
                     detalle, dict(art))


def ajustar_existencia(con, codigo: str, nueva: float, nota: str = "") -> Resultado:
    """Correccion manual, por ejemplo despues de un conteo fisico.

    En un AGRUPADO solo corrige las piezas DISPONIBLES en bodega; lo que
    esta prestado no se toca.
    """
    codigo = codigo.strip().upper()
    art = buscar_articulo(con, codigo)
    if art is None:
        return Resultado(False, "error", f"{codigo} no esta registrado")
    if float(nueva) < 0:
        return Resultado(False, "error", "La existencia no puede ser negativa",
                         "", dict(art))

    diferencia = float(nueva) - float(art["existencia"])
    con.execute(
        "UPDATE articulos SET existencia = ?, estado = ? WHERE codigo = ?",
        (float(nueva), "EN_BODEGA" if nueva > 0 else "FUERA", codigo),
    )
    _guardar_movimiento(con, codigo, None, "AJUSTE", diferencia,
                        nota or "ajuste manual")
    con.commit()
    detalle = f"Existencia {fmt(art['existencia'])} -> {fmt(nueva)}"
    if art["tipo"] == "AGRUPADO":
        prestado = _prestado(art)
        detalle = (f"Disponibles {fmt(art['existencia'])} -> {fmt(nueva)}  |  "
                   f"prestadas {fmt(prestado)} (sin cambio), total "
                   f"{fmt(float(nueva) + prestado)}")
    return Resultado(True, "ok", f"Ajuste: {art['nombre']}", detalle, dict(art))


def _nombre_persona(con, codigo) -> str:
    if not codigo:
        return ""
    fila = con.execute("SELECT nombre FROM personas WHERE codigo = ?",
                       (codigo,)).fetchone()
    return fila["nombre"] if fila else codigo


def _quien_lo_tiene(con, codigo) -> str:
    fila = con.execute(
        "SELECT m.fecha, COALESCE(p.nombre, '(sin operador)') AS quien "
        "FROM movimientos m LEFT JOIN personas p ON p.codigo = m.persona "
        "WHERE m.codigo = ? AND m.tipo = 'SALIDA' ORDER BY m.id DESC LIMIT 1",
        (codigo,),
    ).fetchone()
    return f"Lo saco {fila['quien']} el {fila['fecha']}" if fila else ""


def _quienes_lo_tienen(con, codigo) -> str:
    """'Emerson Garcia (2), Victor Lopez (1)' para un AGRUPADO."""
    return ", ".join(f"{p['nombre']} ({fmt(p['cantidad'])})"
                     for p in prestamos_de(con, codigo))


# --------------------------------------------------------------------------
# Exportacion
# --------------------------------------------------------------------------

def exportar_csv(con: sqlite3.Connection, carpeta: Path) -> list[Path]:
    carpeta = Path(carpeta)
    carpeta.mkdir(parents=True, exist_ok=True)
    generados = []

    consultas = {
        "articulos.csv": "SELECT * FROM articulos ORDER BY codigo",
        "movimientos.csv": (
            "SELECT m.id, m.fecha, m.codigo, a.nombre, m.tipo, m.cantidad, "
            "a.unidad, COALESCE(p.nombre, '') AS persona, m.nota "
            "FROM movimientos m "
            "LEFT JOIN articulos a ON a.codigo = m.codigo "
            "LEFT JOIN personas p ON p.codigo = m.persona ORDER BY m.id"
        ),
        "personas.csv": "SELECT * FROM personas ORDER BY codigo",
    }

    for nombre, sql in consultas.items():
        filas = con.execute(sql).fetchall()
        ruta = carpeta / nombre
        with open(ruta, "w", newline="", encoding="utf-8-sig") as archivo:
            escritor = csv.writer(archivo)
            if filas:
                escritor.writerow(filas[0].keys())
                escritor.writerows([tuple(f) for f in filas])
        generados.append(ruta)

    # Prestamos activos de herramientas con cantidad (quien tiene cuantas).
    ruta = carpeta / "prestamos.csv"
    with open(ruta, "w", newline="", encoding="utf-8-sig") as archivo:
        escritor = csv.writer(archivo)
        escritor.writerow(COLUMNAS_PRESTAMOS)
        escritor.writerows([tuple(f) for f in prestamos_activos(con)])
    generados.append(ruta)

    return generados


COLUMNAS_PRESTAMOS = ("codigo", "articulo", "persona", "nombre_persona",
                      "cantidad", "unidad", "valor_unitario", "valor", "desde")

def exportar_excel_nativo(con: sqlite3.Connection, carpeta: Path) -> Path:
    """Exporta todas las tablas a un solo archivo Excel con pestañas.

    pandas y openpyxl se importan aqui adentro (no al inicio del archivo)
    a proposito: son librerias pesadas que solo hacen falta para esto, y asi
    programas que solo leen (como visor.py) no necesitan instalarlas.
    """
    import pandas as pd
    import openpyxl
    from openpyxl.styles import PatternFill

    carpeta = Path(carpeta)
    carpeta.mkdir(parents=True, exist_ok=True)
    
    from datetime import datetime
    marca = datetime.now().strftime("%Y%m%d")
    ruta_salida = carpeta / f"Reporte_Inventario_{marca}.xlsx"
    # Si está abierto en Excel, esta línea lanzará un PermissionError que la app va a atrapar.
    if ruta_salida.exists():
        ruta_salida.unlink()

    df_articulos = pd.read_sql_query("SELECT * FROM articulos ORDER BY codigo", con)
    df_personas = pd.read_sql_query("SELECT * FROM personas ORDER BY codigo", con)
    df_movimientos = pd.read_sql_query("""
        SELECT m.id, m.fecha, m.codigo, a.nombre, m.tipo, m.cantidad, 
               a.unidad, COALESCE(p.nombre, '') AS persona, m.nota 
        FROM movimientos m 
        LEFT JOIN articulos a ON a.codigo = m.codigo 
        LEFT JOIN personas p ON p.codigo = m.persona 
        ORDER BY m.id DESC
    """, con)
    df_prestamos = pd.DataFrame([tuple(f) for f in prestamos_activos(con)],
                                columns=list(COLUMNAS_PRESTAMOS))

    # Definir los colores (verde para EN_BODEGA, rojo para FUERA)
    fondo_verde = PatternFill(start_color='A9D08E', end_color='A9D08E', fill_type='solid')
    fondo_rojo = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')

    with pd.ExcelWriter(ruta_salida, engine='openpyxl') as writer:
        df_articulos.to_excel(writer, sheet_name='Artículos', index=False)
        df_movimientos.to_excel(writer, sheet_name='Movimientos', index=False)
        df_personas.to_excel(writer, sheet_name='Operadores', index=False)
        df_prestamos.to_excel(writer, sheet_name='Préstamos', index=False)

        for sheetname, df in zip(['Artículos', 'Movimientos', 'Operadores', 'Préstamos'],
                                 [df_articulos, df_movimientos, df_personas, df_prestamos]):
            worksheet = writer.sheets[sheetname]
            
            # Auto-ajustar el ancho de las columnas
            for i, col in enumerate(df.columns):
                largo_datos = df[col].astype(str).map(len).max() if len(df) else 0
                max_len = max(largo_datos, len(col)) + 2
                worksheet.column_dimensions[openpyxl.utils.get_column_letter(i+1)].width = max_len

            # Aplicar colores en la pestaña Artículos
            if sheetname == 'Artículos' and 'estado' in df.columns:
                # Obtener el índice numérico de la columna 'estado'
                col_estado_idx = df.columns.get_loc('estado') + 1
                
                # Iterar sobre las filas (saltando la fila 1 que es el encabezado)
                for row_idx in range(2, len(df) + 2):
                    celda = worksheet.cell(row=row_idx, column=col_estado_idx)
                    if celda.value == 'EN_BODEGA':
                        celda.fill = fondo_verde
                    elif celda.value == 'FUERA':
                        celda.fill = fondo_rojo

    return ruta_salida